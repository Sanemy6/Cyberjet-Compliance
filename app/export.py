"""Export des questionnaires traités.

Deux modes :
  * « original » — réinjection des réponses dans le classeur Excel reçu du
    client, à l'emplacement exact repéré à l'import. C'est ce que le client
    attend : son propre fichier, complété.
  * « rapport »  — classeur récapitulatif autonome, disponible quel que soit
    le format d'origine (Word, PDF, saisie manuelle).
"""

from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .db import EXPORT_DIR
from .search import band

_HEADER_FILL = PatternFill("solid", fgColor="0F2740")
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
_THIN = Side(style="thin", color="D6DEE8")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_WRAP = Alignment(wrap_text=True, vertical="top")

_BAND_FILL = {
    "haute": PatternFill("solid", fgColor="E4F5EC"),
    "moyenne": PatternFill("solid", fgColor="FFF4E0"),
    "faible": PatternFill("solid", fgColor="FDE9E7"),
    "aucune": PatternFill("solid", fgColor="F1F3F7"),
}

_SAFE = re.compile(r"[^A-Za-z0-9_.\- ]+")


def _safe_name(name: str) -> str:
    return _SAFE.sub("_", name).strip() or "export"


def _answer_of(question: dict) -> str:
    return (question.get("final_answer") or question.get("suggested_answer") or "").strip()


def export_original(questionnaire: dict, questions: list[dict]) -> Path:
    """Recopie les réponses dans le classeur d'origine et enregistre une copie."""
    source = Path(questionnaire["stored_path"])
    if not source.exists():
        raise FileNotFoundError("Le fichier d'origine est introuvable.")

    wb = load_workbook(source)
    written = 0
    headers_done: set[tuple[str, int]] = set()

    for question in questions:
        locator = question.get("locator_json")
        if not locator:
            continue
        loc = json.loads(locator) if isinstance(locator, str) else locator
        sheet = loc.get("sheet")
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        answer = _answer_of(question)
        if not answer:
            continue

        col = int(loc["answer_col"])
        row = int(loc["row"])

        if loc.get("answer_created"):
            key = (sheet, col)
            if key not in headers_done:
                header_row = int(loc.get("header_row") or 1)
                if header_row >= 1:
                    cell = ws.cell(row=header_row, column=col, value="Réponse")
                    cell.fill = _HEADER_FILL
                    cell.font = _HEADER_FONT
                    cell.alignment = _WRAP
                ws.column_dimensions[get_column_letter(col)].width = 70
                headers_done.add(key)

        cell = ws.cell(row=row, column=col, value=answer)
        cell.alignment = _WRAP
        written += 1

    stem = _safe_name(Path(questionnaire.get("original_filename") or questionnaire["name"]).stem)
    stamp = datetime.now().strftime("%Y%m%d-%H%M")
    target = EXPORT_DIR / f"{stem}_complete_{stamp}.xlsx"
    wb.save(target)
    wb.close()

    if not written:
        raise ValueError("Aucune réponse à écrire : validez au moins une question.")
    return target


def export_report(questionnaire: dict, questions: list[dict], client_name: str) -> Path:
    """Classeur récapitulatif : synthèse + détail question par question."""
    wb = Workbook()

    summary = wb.active
    summary.title = "Synthèse"
    total = len(questions)
    validated = sum(1 for q in questions if q["status"] == "validated")
    na = sum(1 for q in questions if q["status"] == "na")
    answered = sum(1 for q in questions if _answer_of(q))

    meta = [
        ("Client", client_name),
        ("Questionnaire", questionnaire["name"]),
        ("Fichier d'origine", questionnaire.get("original_filename") or "—"),
        ("Date d'export", datetime.now().strftime("%d/%m/%Y %H:%M")),
        ("", ""),
        ("Questions", total),
        ("Réponses proposées", answered),
        ("Réponses validées", validated),
        ("Sans objet", na),
        ("Taux de couverture", f"{(answered / total * 100):.0f} %" if total else "—"),
        ("Taux de validation", f"{(validated / total * 100):.0f} %" if total else "—"),
    ]
    for r, (label, value) in enumerate(meta, start=1):
        summary.cell(row=r, column=1, value=label).font = Font(bold=True)
        summary.cell(row=r, column=2, value=value)
    summary.column_dimensions["A"].width = 26
    summary.column_dimensions["B"].width = 60

    detail = wb.create_sheet("Réponses")
    headers = ["N°", "Section", "Réf.", "Question", "Réponse", "Confiance",
               "Statut", "Source"]
    for c, title in enumerate(headers, start=1):
        cell = detail.cell(row=1, column=c, value=title)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _WRAP
        cell.border = _BORDER

    status_label = {
        "validated": "Validée",
        "auto": "Proposée",
        "pending": "À traiter",
        "na": "Sans objet",
    }

    for r, question in enumerate(questions, start=2):
        confidence = float(question.get("confidence") or 0)
        level = band(confidence)
        values = [
            question["idx"],
            question.get("section") or "",
            question.get("ref") or "",
            question["text"],
            _answer_of(question),
            f"{confidence * 100:.0f} %" if confidence else "—",
            status_label.get(question["status"], question["status"]),
            _source_label(question),
        ]
        for c, value in enumerate(values, start=1):
            cell = detail.cell(row=r, column=c, value=value)
            cell.alignment = _WRAP
            cell.border = _BORDER
        detail.cell(row=r, column=6).fill = _BAND_FILL.get(level, _BAND_FILL["aucune"])

    for column, width in zip("ABCDEFGH", (6, 22, 14, 60, 70, 12, 14, 28)):
        detail.column_dimensions[column].width = width
    detail.freeze_panes = "A2"

    stem = _safe_name(questionnaire["name"])
    stamp = datetime.now().strftime("%Y%m%d-%H%M")
    target = EXPORT_DIR / f"{stem}_rapport_{stamp}.xlsx"
    wb.save(target)
    wb.close()
    return target


def _source_label(question: dict) -> str:
    kind = question.get("matched_kind")
    if kind == "kb":
        return "Base de connaissance"
    if kind == "config":
        return "Configuration client"
    return "Saisie manuelle" if _answer_of(question) else ""
