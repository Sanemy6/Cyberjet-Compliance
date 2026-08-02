"""Extraction des questions depuis un questionnaire client (xlsx / docx / pdf).

L'enjeu n'est pas de lire le fichier — c'est de deviner sa structure, car
chaque client a la sienne. La stratégie :

  * Excel : on cherche la ligne d'en-tête, puis on identifie par mots-clés les
    colonnes Question / Réponse / Référence / Section. En dernier recours, la
    colonne « question » est celle dont le texte est le plus long en moyenne.
    Les coordonnées (feuille, ligne, colonne) sont conservées : elles
    permettent de réinjecter les réponses dans le classeur d'origine.

  * Word : les tableaux d'abord (structure la plus fréquente), puis les
    paragraphes numérotés qui ressemblent à des questions.

  * PDF : les tableaux détectés par pdfplumber, puis le texte découpé en items
    numérotés. Pas de réinjection possible — l'export se fait en Excel.
"""

from __future__ import annotations

import re
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any

from .textproc import looks_like_question, normalize, strip_accents

# --------------------------------------------------------------------------
# Mots-clés d'en-tête (comparés sur la forme normalisée, sans accents)
# --------------------------------------------------------------------------
QUESTION_HEADERS = (
    "question", "questions", "exigence", "requirement", "control objective",
    "controle", "control", "critere", "criteria", "description", "libelle",
    "intitule", "enonce", "demande", "objet", "item", "mesure", "point de controle",
    "expected", "attendu", "sujet", "ask",
)
ANSWER_HEADERS = (
    "reponse", "answer", "response", "commentaire", "comment", "remarque",
    "remark", "statut", "status", "compliance", "conformite", "conforme",
    "justification", "explication", "explanation", "detail", "observation",
    "supplier response", "vendor response", "reponse fournisseur",
)
REF_HEADERS = (
    "ref", "reference", "id", "n", "no", "num", "numero", "code", "article",
    "clause", "question id", "item id", "identifiant", "index",
)
SECTION_HEADERS = (
    "section", "domaine", "domain", "categorie", "category", "theme", "chapitre",
    "chapter", "famille", "groupe", "group", "area", "topic", "pilier", "volet",
)
EVIDENCE_HEADERS = ("preuve", "evidence", "piece", "attachment", "document", "proof")

_NUMBERING = re.compile(
    r"^\s*(?:"
    r"(?:Q|R|C|A)?\d+(?:[.\-]\d+)*"        # 1  /  1.2.3  /  Q12  /  A1-2
    r"|[A-Z]\.\d+(?:\.\d+)*"               # A.8.24
    r"|\(?[a-z]\)"                         # (a)
    r")[\.\)\-–:]?\s+",
    re.I,
)


@dataclass
class ExtractedQuestion:
    idx: int
    text: str
    ref: str | None = None
    section: str | None = None
    locator: dict[str, Any] | None = None


@dataclass
class ParseResult:
    questions: list[ExtractedQuestion] = field(default_factory=list)
    file_type: str = ""
    warnings: list[str] = field(default_factory=list)
    layout: dict[str, Any] = field(default_factory=dict)

    def to_dict(self) -> dict[str, Any]:
        return {
            "questions": [asdict(q) for q in self.questions],
            "file_type": self.file_type,
            "warnings": self.warnings,
            "layout": self.layout,
        }


def detect_file_type(filename: str) -> str:
    ext = Path(filename).suffix.lower().lstrip(".")
    if ext in ("xlsx", "xlsm", "xls"):
        return "xlsx"
    if ext in ("docx", "doc"):
        return "docx"
    if ext == "pdf":
        return "pdf"
    if ext in ("csv", "txt"):
        return "csv"
    return ext or "inconnu"


def parse(path: Path, filename: str | None = None) -> ParseResult:
    kind = detect_file_type(filename or path.name)
    if kind == "xlsx":
        return parse_excel(path)
    if kind == "docx":
        return parse_word(path)
    if kind == "pdf":
        return parse_pdf(path)
    if kind == "csv":
        return parse_csv(path)
    return ParseResult(file_type=kind, warnings=[f"Format « {kind} » non pris en charge."])


# --------------------------------------------------------------------------
# Utilitaires communs
# --------------------------------------------------------------------------

def _clean(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r", " ").strip()
    return re.sub(r"[ \t]+", " ", text)


def _header_score(text: str, keywords: tuple[str, ...]) -> int:
    """Score de correspondance d'un en-tête : exact > préfixe > inclusion."""
    flat = normalize(text)
    if not flat:
        return 0
    best = 0
    for kw in keywords:
        if flat == kw:
            best = max(best, 3)
        elif flat.startswith(kw) or flat.endswith(kw):
            best = max(best, 2)
        elif kw in flat and len(kw) > 3:
            best = max(best, 1)
    return best


def _strip_numbering(text: str) -> tuple[str | None, str]:
    """Sépare un éventuel préfixe de numérotation du corps de la question."""
    match = _NUMBERING.match(text)
    if not match:
        return None, text
    ref = match.group(0).strip(" .)-–:")
    body = text[match.end():].strip()
    # On ne coupe que si ce qui reste est substantiel.
    if len(body) < 10:
        return None, text
    return ref or None, body


# --------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------

def parse_excel(path: Path) -> ParseResult:
    from openpyxl import load_workbook

    result = ParseResult(file_type="xlsx")
    wb = load_workbook(path, data_only=True, read_only=False)
    idx = 0
    sheets_layout: list[dict[str, Any]] = []

    for ws in wb.worksheets:
        if ws.sheet_state != "visible" or ws.max_row < 2:
            continue

        merged = _merged_map(ws)
        grid = _read_grid(ws, merged, max_rows=min(ws.max_row, 4000))
        if not grid:
            continue

        header_row, columns = _detect_columns(grid)
        if columns.get("question") is None:
            continue

        qcol = columns["question"]
        acol = columns.get("answer")
        if acol is None:
            acol = _first_free_column(grid, header_row, qcol)
            columns["answer"] = acol
            columns["answer_created"] = True

        sheets_layout.append({
            "sheet": ws.title,
            "header_row": header_row + 1,
            **{k: (v + 1 if isinstance(v, int) and k.endswith("col") else v)
               for k, v in _named_columns(columns).items()},
        })

        current_section = None
        for r in range(header_row + 1, len(grid)):
            row = grid[r]
            question = _clean(row[qcol]) if qcol < len(row) else ""
            if not question:
                continue

            filled = sum(1 for c in row if _clean(c))
            # Ligne ne contenant qu'un libellé court : c'est un titre de section.
            if filled == 1 and len(question) < 70 and not question.endswith("?"):
                current_section = question
                continue
            if len(question) < 10:
                continue

            ref = None
            if columns.get("ref") is not None and columns["ref"] < len(row):
                ref = _clean(row[columns["ref"]]) or None
            section = current_section
            if columns.get("section") is not None and columns["section"] < len(row):
                section = _clean(row[columns["section"]]) or current_section

            if ref is None:
                ref, question = _strip_numbering(question)

            idx += 1
            result.questions.append(ExtractedQuestion(
                idx=idx,
                text=question,
                ref=ref,
                section=section,
                locator={
                    "sheet": ws.title,
                    "row": r + 1,            # 1-indexé, comme openpyxl
                    "question_col": qcol + 1,
                    "answer_col": acol + 1,
                    "header_row": header_row + 1,
                    "answer_created": columns.get("answer_created", False),
                },
            ))

    wb.close()
    result.layout = {"sheets": sheets_layout}
    if not result.questions:
        result.warnings.append(
            "Aucune colonne de questions identifiée. Vérifiez que le classeur "
            "contient bien une colonne d'intitulés (Question, Exigence, Contrôle…)."
        )
    return result


def _merged_map(ws) -> dict[tuple[int, int], tuple[int, int]]:
    """(ligne, col) -> (ligne, col) de la cellule porteuse de la valeur fusionnée."""
    mapping: dict[tuple[int, int], tuple[int, int]] = {}
    for rng in ws.merged_cells.ranges:
        anchor = (rng.min_row, rng.min_col)
        for row in range(rng.min_row, rng.max_row + 1):
            for col in range(rng.min_col, rng.max_col + 1):
                mapping[(row, col)] = anchor
    return mapping


def _read_grid(ws, merged, max_rows: int) -> list[list[Any]]:
    grid: list[list[Any]] = []
    for r, row in enumerate(ws.iter_rows(min_row=1, max_row=max_rows), start=1):
        values: list[Any] = []
        for c, cell in enumerate(row, start=1):
            value = cell.value
            if value is None and (r, c) in merged:
                ar, ac = merged[(r, c)]
                value = ws.cell(row=ar, column=ac).value
            values.append(value)
        grid.append(values)
    return grid


def _detect_columns(grid: list[list[Any]]) -> tuple[int, dict[str, Any]]:
    """Localise la ligne d'en-tête et le rôle de chaque colonne."""
    scan = min(len(grid), 25)
    best_row, best_score = 0, -1
    for r in range(scan):
        score = 0
        for cell in grid[r]:
            text = _clean(cell)
            if not text or len(text) > 60:
                continue
            score += max(
                _header_score(text, QUESTION_HEADERS),
                _header_score(text, ANSWER_HEADERS),
                _header_score(text, REF_HEADERS),
                _header_score(text, SECTION_HEADERS),
            )
        if score > best_score:
            best_row, best_score = r, score

    header = grid[best_row] if best_score > 0 else []
    columns: dict[str, Any] = {}
    for role, keywords in (
        ("question", QUESTION_HEADERS),
        ("answer", ANSWER_HEADERS),
        ("ref", REF_HEADERS),
        ("section", SECTION_HEADERS),
        ("evidence", EVIDENCE_HEADERS),
    ):
        best_col, top = None, 0
        for c, cell in enumerate(header):
            text = _clean(cell)
            if not text or len(text) > 60:
                continue
            s = _header_score(text, keywords)
            if s > top:
                best_col, top = c, s
        columns[role] = best_col

    # Repli : la colonne « question » est celle au texte le plus long.
    if columns.get("question") is None:
        columns["question"] = _widest_text_column(grid, best_row)
        if best_score <= 0:
            best_row = -1  # pas d'en-tête : les données commencent en ligne 1

    # Une même colonne ne peut pas être à la fois question et réponse.
    if columns.get("answer") == columns.get("question"):
        columns["answer"] = None
    return best_row, columns


def _widest_text_column(grid: list[list[Any]], header_row: int) -> int | None:
    width = max((len(r) for r in grid), default=0)
    best_col, best_avg = None, 0.0
    for c in range(width):
        lengths = [
            len(_clean(grid[r][c]))
            for r in range(header_row + 1, len(grid))
            if c < len(grid[r]) and _clean(grid[r][c])
        ]
        if len(lengths) < 3:
            continue
        avg = sum(lengths) / len(lengths)
        if avg > best_avg and avg >= 20:
            best_col, best_avg = c, avg
    return best_col


def _first_free_column(grid: list[list[Any]], header_row: int, qcol: int) -> int:
    """Première colonne entièrement vide après la colonne des questions."""
    width = max((len(r) for r in grid), default=0)
    for c in range(qcol + 1, width + 2):
        empty = all(
            c >= len(grid[r]) or not _clean(grid[r][c])
            for r in range(max(header_row, 0), len(grid))
        )
        if empty:
            return c
    return width


def _named_columns(columns: dict[str, Any]) -> dict[str, Any]:
    return {f"{role}_col": value for role, value in columns.items()
            if role != "answer_created"}


# --------------------------------------------------------------------------
# Word
# --------------------------------------------------------------------------

def parse_word(path: Path) -> ParseResult:
    import docx

    result = ParseResult(file_type="docx")
    document = docx.Document(str(path))
    idx = 0

    for t, table in enumerate(document.tables):
        rows = [[_clean(cell.text) for cell in row.cells] for row in table.rows]
        rows = [r for r in rows if any(r)]
        if len(rows) < 2:
            continue

        header = rows[0]
        qcol, top = None, 0
        for c, cell in enumerate(header):
            s = _header_score(cell, QUESTION_HEADERS)
            if s > top:
                qcol, top = c, s
        start = 1 if top > 0 else 0
        if qcol is None:
            # Pas d'en-tête exploitable : colonne au texte le plus long.
            widths = [
                (sum(len(r[c]) for r in rows if c < len(r)) / max(len(rows), 1), c)
                for c in range(max(len(r) for r in rows))
            ]
            qcol = max(widths)[1] if widths else 0
            start = 0

        for r in range(start, len(rows)):
            row = rows[r]
            if qcol >= len(row):
                continue
            text = row[qcol]
            if len(text) < 12:
                continue
            ref, text = _strip_numbering(text)
            idx += 1
            result.questions.append(ExtractedQuestion(
                idx=idx, text=text, ref=ref, section=f"Tableau {t + 1}"
            ))

    # Paragraphes : uniquement ceux qui ressemblent vraiment à une question.
    section = None
    for para in document.paragraphs:
        text = _clean(para.text)
        if not text:
            continue
        style = (para.style.name or "").lower()
        if style.startswith("heading") or style.startswith("titre"):
            section = text
            continue
        if not looks_like_question(text):
            continue
        ref, body = _strip_numbering(text)
        idx += 1
        result.questions.append(ExtractedQuestion(
            idx=idx, text=body, ref=ref, section=section
        ))

    if not result.questions:
        result.warnings.append(
            "Aucune question détectée dans le document Word. "
            "Vous pouvez saisir les questions manuellement."
        )
    return result


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------

def parse_pdf(path: Path) -> ParseResult:
    import pdfplumber

    result = ParseResult(file_type="pdf")
    idx = 0
    seen: set[str] = set()

    with pdfplumber.open(str(path)) as pdf:
        for p, page in enumerate(pdf.pages, start=1):
            handled_rows = 0
            try:
                tables = page.extract_tables() or []
            except Exception:
                tables = []

            for table in tables:
                rows = [[_clean(c) for c in (row or [])] for row in table]
                rows = [r for r in rows if any(r)]
                if len(rows) < 2 or len(rows[0]) < 2:
                    continue
                qcol, top = 0, 0
                for c, cell in enumerate(rows[0]):
                    s = _header_score(cell, QUESTION_HEADERS)
                    if s > top:
                        qcol, top = c, s
                if top == 0:
                    lengths = [
                        sum(len(r[c]) for r in rows if c < len(r))
                        for c in range(len(rows[0]))
                    ]
                    qcol = lengths.index(max(lengths))
                for row in rows[1 if top else 0:]:
                    if qcol >= len(row):
                        continue
                    text = row[qcol]
                    if len(text) < 12:
                        continue
                    key = normalize(text)[:120]
                    if key in seen:
                        continue
                    seen.add(key)
                    ref, body = _strip_numbering(text)
                    idx += 1
                    handled_rows += 1
                    result.questions.append(ExtractedQuestion(
                        idx=idx, text=body, ref=ref, section=f"Page {p}"
                    ))

            if handled_rows:
                continue

            text = page.extract_text() or ""
            for block in _pdf_blocks(text):
                if not looks_like_question(block):
                    continue
                key = normalize(block)[:120]
                if key in seen:
                    continue
                seen.add(key)
                ref, body = _strip_numbering(block)
                idx += 1
                result.questions.append(ExtractedQuestion(
                    idx=idx, text=body, ref=ref, section=f"Page {p}"
                ))

    if not result.questions:
        result.warnings.append(
            "Aucune question détectée dans le PDF (document scanné ou mise en page "
            "atypique). Vous pouvez saisir les questions manuellement."
        )
    return result


def _pdf_blocks(text: str) -> list[str]:
    """Regroupe les lignes en blocs : un nouvel item numéroté ouvre un bloc."""
    blocks: list[str] = []
    current: list[str] = []
    for raw in text.splitlines():
        line = _clean(raw)
        if not line:
            if current:
                blocks.append(" ".join(current))
                current = []
            continue
        starts_item = bool(_NUMBERING.match(line)) or line.endswith("?")
        if starts_item and current:
            blocks.append(" ".join(current))
            current = []
        current.append(line)
        if line.endswith("?") and current:
            blocks.append(" ".join(current))
            current = []
    if current:
        blocks.append(" ".join(current))
    return [b for b in blocks if len(b) >= 12]


# --------------------------------------------------------------------------
# CSV / texte brut
# --------------------------------------------------------------------------

def parse_csv(path: Path) -> ParseResult:
    import csv
    import io

    result = ParseResult(file_type="csv")
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = raw.decode("utf-8", errors="replace")

    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
    except csv.Error:
        dialect = csv.excel
        dialect.delimiter = ";" if sample.count(";") > sample.count(",") else ","

    rows = [r for r in csv.reader(io.StringIO(text), dialect) if any(_clean(c) for c in r)]
    if not rows:
        result.warnings.append("Fichier vide.")
        return result

    header = [_clean(c) for c in rows[0]]
    qcol, top = 0, 0
    for c, cell in enumerate(header):
        s = _header_score(cell, QUESTION_HEADERS)
        if s > top:
            qcol, top = c, s
    if top == 0 and len(header) == 1:
        qcol = 0

    idx = 0
    for row in rows[1 if top else 0:]:
        if qcol >= len(row):
            continue
        text_value = _clean(row[qcol])
        if len(text_value) < 10:
            continue
        ref, body = _strip_numbering(text_value)
        idx += 1
        result.questions.append(ExtractedQuestion(idx=idx, text=body, ref=ref))
    return result


def parse_text_block(raw: str) -> ParseResult:
    """Saisie manuelle : une question par ligne (ou séparées par une ligne vide)."""
    result = ParseResult(file_type="manuel")
    idx = 0
    for line in raw.splitlines():
        text = _clean(line)
        if len(text) < 8:
            continue
        ref, body = _strip_numbering(text)
        idx += 1
        result.questions.append(ExtractedQuestion(idx=idx, text=body, ref=ref))
    return result
